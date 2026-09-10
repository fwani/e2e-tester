"""내보낸 파일의 서식 (014 T075 · US4).

이 파일은 **회의 자료가 된다.** 통과와 부분 통과가 같아 보이면 보고가 틀리고, 머리글이
스크롤에 사라지면 아래쪽 행을 읽을 수 없다.

**색만으로 구분하지 않는다** — 칸 안에 이미 글자가 있고 색은 거드는 것뿐이다. 그래서
여기서 확인하는 것은 "색이 서로 다른가"이지 "무슨 색인가"가 아니다.
"""

from __future__ import annotations

import io

from excel_support import make_result, make_test, repo_of
from fastapi.testclient import TestClient

from itb.domain.run_result import Outcome
from itb.portability.sheet_name import UNGROUPED_SHEET_NAME

EXPORT = "/api/export"


def opened(data: bytes):
    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(data))


def seed_outcomes(client: TestClient) -> None:
    repo = repo_of(client)
    for i, outcome in enumerate(
        [Outcome.PASS, Outcome.FAIL, Outcome.PARTIAL_PASS, Outcome.STOPPED], start=1
    ):
        tid = f"TC-{i:03d}"
        repo.write_test(make_test(tid, f"테스트 {i}"))
        repo.write_result(make_result(tid, outcome))
    # 미실행 하나를 더한다 — 빈 칸이 어떤 서식도 받지 않는지 본다.
    repo.write_test(make_test("TC-005", "안 돌린 것"))


class HeaderTests:
    def test_머리글_행이_고정된다(self, project_client: TestClient) -> None:
        wb = opened(project_client.get(EXPORT).content)
        assert wb[UNGROUPED_SHEET_NAME].freeze_panes == "A2"

    def test_머리글이_굵고_배경이_있다(self, project_client: TestClient) -> None:
        ws = opened(project_client.get(EXPORT).content)[UNGROUPED_SHEET_NAME]
        assert ws["A1"].font.bold
        assert ws["A1"].fill.fgColor.rgb not in (None, "00000000")

    def test_모든_시트에_걸린다(self, project_client: TestClient) -> None:
        project_client.post("/api/groups", json={"prefix": "USER", "name": "사용자관리"})
        wb = opened(project_client.get(EXPORT).content)
        for name in wb.sheetnames:
            assert wb[name].freeze_panes == "A2"


class ColumnWidthTests:
    def test_모든_열에_너비가_잡힌다(self, project_client: TestClient) -> None:
        ws = opened(project_client.get(EXPORT).content)[UNGROUPED_SHEET_NAME]
        widths = [ws.column_dimensions[c].width for c in "ABCDEFG"]
        assert all(w and w > 0 for w in widths)

    def test_절차와_기대결과가_더_넓다(self, project_client: TestClient) -> None:
        # 여러 줄이 들어가는 칸이다. 좁으면 읽을 수 없다.
        ws = opened(project_client.get(EXPORT).content)[UNGROUPED_SHEET_NAME]
        assert ws.column_dimensions["E"].width > ws.column_dimensions["A"].width
        assert ws.column_dimensions["F"].width > ws.column_dimensions["A"].width

    def test_여러_줄_칸이_줄바꿈된다(self, project_client: TestClient) -> None:
        repo_of(project_client).write_test(make_test("TC-001", "여러 줄"))
        ws = opened(project_client.get(EXPORT).content)[UNGROUPED_SHEET_NAME]
        assert ws["E2"].alignment.wrap_text


class OutcomeStyleTests:
    def test_네_결말이_서로_다른_서식을_받는다(self, project_client: TestClient) -> None:
        seed_outcomes(project_client)
        ws = opened(project_client.get(EXPORT).content)[UNGROUPED_SHEET_NAME]
        fills = [ws[f"G{r}"].fill.fgColor.rgb for r in range(2, 6)]
        assert len(set(fills)) == 4

    def test_통과와_실패가_다르게_보인다(self, project_client: TestClient) -> None:
        seed_outcomes(project_client)
        ws = opened(project_client.get(EXPORT).content)[UNGROUPED_SHEET_NAME]
        assert ws["G2"].value == "P"
        assert ws["G3"].value == "F"
        assert ws["G2"].fill.fgColor.rgb != ws["G3"].fill.fgColor.rgb

    def test_부분_통과가_통과와_다르게_보인다(self, project_client: TestClient) -> None:
        # 보고서에서 이 둘이 같아 보이면 오보다.
        seed_outcomes(project_client)
        ws = opened(project_client.get(EXPORT).content)[UNGROUPED_SHEET_NAME]
        assert ws["G4"].value == "P(부분)"
        assert ws["G4"].fill.fgColor.rgb != ws["G2"].fill.fgColor.rgb

    def test_미실행은_어떤_서식도_받지_않는다(self, project_client: TestClient) -> None:
        seed_outcomes(project_client)
        ws = opened(project_client.get(EXPORT).content)[UNGROUPED_SHEET_NAME]
        assert ws["G6"].value in (None, "")
        assert ws["G6"].fill.fgColor.rgb in (None, "00000000")

    def test_색_없이도_읽힌다(self, project_client: TestClient) -> None:
        # 색만으로 구분하지 않는다 — 칸 안에 글자가 있다.
        seed_outcomes(project_client)
        ws = opened(project_client.get(EXPORT).content)[UNGROUPED_SHEET_NAME]
        assert [ws[f"G{r}"].value for r in range(2, 6)] == ["P", "F", "P(부분)", "중지"]

    def test_결과_칸이_가운데_정렬된다(self, project_client: TestClient) -> None:
        seed_outcomes(project_client)
        ws = opened(project_client.get(EXPORT).content)[UNGROUPED_SHEET_NAME]
        assert ws["G2"].alignment.horizontal == "center"
