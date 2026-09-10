"""엑셀 통로 검증용 공용 도구 (기능 014).

브라우저 없이 프로젝트를 채운다. 내보내기가 보는 것은 저장된 정의와 실행 결과뿐이므로,
그 둘을 직접 써 넣으면 녹화를 돌리지 않고도 모든 모양을 만들 수 있다.
"""

from __future__ import annotations

import datetime as dt
import io
import pathlib
from typing import Any

from fastapi.testclient import TestClient

from itb.domain.assertion import Assertion, AssertionKind
from itb.domain.locator import Candidate, CandidateStatus, TargetLocator
from itb.domain.run_result import Outcome, RunResult, RunScope
from itb.domain.step import AssertionStep, ClickStep, FillStep, Step
from itb.domain.test_case import AuthoringMode, Test, Variable
from itb.storage.repository import ProjectRepository


def verified(value: str) -> Candidate:
    return Candidate(value=value, status=CandidateStatus.VERIFIED)


def target(text: str) -> TargetLocator:
    return TargetLocator(text=verified(text))


def click(index: int, label: str) -> Step:
    return ClickStep(id=f"step-{index:02d}", label=label, target=target(label))


def fill(index: int, label: str, value: str) -> Step:
    return FillStep(id=f"step-{index:02d}", label=label, target=target(label), value=value)


def check(index: int, label: str, value: str) -> Step:
    return AssertionStep(
        id=f"step-{index:02d}",
        label=label,
        assertion=Assertion(kind=AssertionKind.TEXT, value=value),
    )


def make_test(
    test_id: str,
    name: str,
    *,
    steps: list[Step] | None = None,
    variables: list[Variable] | None = None,
    description: str | None = None,
    actor: str | None = None,
) -> Test:
    """저장 가능한 테스트 하나. 기본은 동작 2개 + 검증 1개."""
    return Test(
        id=test_id,
        name=name,
        description=description,
        actor=actor,
        authoring_mode=AuthoringMode.RECORD,
        start_url="https://example.internal/login",
        variables=variables or [],
        steps=steps
        or [
            click(1, "로그인 버튼"),
            click(2, "확인 버튼"),
            check(3, "대시보드 표시 확인", "대시보드"),
        ],
    )


def make_result(test_id: str, outcome: Outcome) -> RunResult:
    now = dt.datetime.now(dt.UTC)
    return RunResult(
        test_id=test_id,
        outcome=outcome,
        started_at=now,
        finished_at=now,
        total_ms=10,
        passed_count=1,
        total_count=1,
        browser="chromium",
        steps=[],
        scope=RunScope.FULL,
        attempted_count=1,
    )


def repo_of(client: TestClient) -> ProjectRepository:
    """열린 프로젝트의 저장소를 연다."""
    root = client.get("/api/project").json()["root"]
    return ProjectRepository.open(pathlib.Path(root))


def read_back(data: bytes) -> dict[str, list[list[Any]]]:
    """내보낸 바이트를 다시 읽어 ``{시트 이름: 행들}`` 로 만든다.

    **openpyxl 로 다시 읽는다.** 우리가 만든 자료구조를 그대로 확인하면 "썼다고 생각한 것"을
    확인하게 된다. 파일을 거쳐 오면 실제로 파일에 들어간 것을 본다.
    """
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return {
            name: [list(row) for row in wb[name].iter_rows(values_only=True)]
            for name in wb.sheetnames
        }
    finally:
        wb.close()


def sheet_names(data: bytes) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()
